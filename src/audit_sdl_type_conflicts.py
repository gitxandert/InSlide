"""Report contrasting SDL Types for the same scanner/date/accession."""

from __future__ import annotations

import argparse
import os
import sys
from zipfile import BadZipFile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from sdl_type_rules import format_conflict, resolve_types, type_record


REQUIRED_HEADERS = ("Accession ID", "Type", "Scanner", "Date Loaded")
DEFAULT_WORKBOOK = Path(__file__).parent / "logs" / "Slide_Digitization_Log.xlsx"


def find_conflicts(path: Path, sheet_name: str):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"worksheet not found: {sheet_name}")
        worksheet = workbook[sheet_name]
        locations: dict[str, list[int]] = {header: [] for header in REQUIRED_HEADERS}
        for column in range(1, worksheet.max_column + 1):
            header = str(worksheet.cell(row=1, column=column).value or "").strip()
            if header in locations:
                locations[header].append(column)
        invalid = [
            header for header, columns in locations.items() if len(columns) != 1
        ]
        if invalid:
            raise ValueError(
                "expected exactly one column for: " + ", ".join(invalid)
            )

        records = []
        for row_number in range(2, worksheet.max_row + 1):
            values = {
                header: worksheet.cell(row=row_number, column=locations[header][0]).value
                for header in REQUIRED_HEADERS
            }
            record = type_record(values, row_number)
            if record is not None:
                records.append(record)
        return resolve_types(records)[1]
    finally:
        workbook.close()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path(os.environ.get("SDL_FILE_PATH", DEFAULT_WORKBOOK)),
        help="SDL workbook path (default: SDL_FILE_PATH or application default)",
    )
    parser.add_argument("--sheet", default="general", help="SDL worksheet name")
    args = parser.parse_args(argv)
    try:
        conflicts = find_conflicts(args.workbook, args.sheet)
    except (OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        print(f"Could not audit SDL workbook: {exc}", file=sys.stderr)
        return 2

    if not conflicts:
        print("No contrasting SDL Types found.")
        return 0
    print(f"Found {len(conflicts)} contrasting SDL Type group(s):")
    for conflict in conflicts:
        print(f"- {format_conflict(conflict)}")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
